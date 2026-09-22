import json
import sys
from pathlib import Path

import openpyxl

SHEET_NAME = "StorageTiers"


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        return []
    header = rows[0]
    data_rows = rows[2:]
    result = []
    for row in data_rows:
        if row is None or all(v is None for v in row):
            continue
        result.append(dict(zip(header, row)))
    return result


def _normalize_value(v):
    if isinstance(v, str):
        v = v.strip()
    return v


def build_storage_tiers(xlsx_path: Path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if SHEET_NAME not in wb.sheetnames:
        raise SystemExit(f"오류: '{xlsx_path}' 파일에 '{SHEET_NAME}' 시트가 없습니다.")

    ws = wb[SHEET_NAME]
    tiers = []
    for row in sheet_rows(ws):
        tier = _normalize_value(row.get("Tier"))
        slot_count = _normalize_value(row.get("SlotCount"))
        cols = _normalize_value(row.get("Cols"))
        if tier is None or slot_count is None or cols is None:
            continue
        tiers.append({
            "tier": int(tier),
            "slotCount": int(slot_count),
            "cols": int(cols),
        })

    tiers.sort(key=lambda t: t["tier"])
    return tiers


def main():
    if len(sys.argv) < 3:
        print("사용법: python StorageTierDataConverter.py <StorageTierData.xlsx> <출력폴더>")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    tiers = build_storage_tiers(xlsx_path)

    out_path = out_dir / "storage_tiers.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(tiers, f, ensure_ascii=False, indent=1)

    print(f"{len(tiers)}개 등급 저장 완료 -> {out_path}")


if __name__ == "__main__":
    main()
