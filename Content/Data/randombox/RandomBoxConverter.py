import csv
import json
import sys
from pathlib import Path

import openpyxl


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        return []
    header = rows[0]
    data_rows = rows[2:]
    result = []
    for row in data_rows:
        if row is None or all(v is None for v in row):
            continue
        result.append(dict(zip(header, row)))
    return result


def _norm(v):
    if isinstance(v, str):
        v = v.strip()
    return v


def _is_true(v):
    return str(v).strip().upper() == "TRUE"


def load_items_json(out_dir: Path):
    items_path = out_dir / "items.json"
    if not items_path.exists():
        raise SystemExit(f"오류: items.json을 찾을 수 없습니다: {items_path} (먼저 ItemDataConverter를 실행하세요)")
    with open(items_path, "r", encoding="utf-8") as f:
        items = json.load(f)
    by_index = {}
    for item in items:
        by_index[item["index"]] = {
            "category": item.get("category", ""),
            "rarityId": item.get("rarityId", ""),
        }
    return by_index


def load_localization(wb):
    # items.xlsx의 Localization 시트와 동일한 구조: 타입 행이 없다 (1행=헤더, 2행부터 데이터).
    if "Localization" not in wb.sheetnames:
        return {}
    ws = wb["Localization"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {}
    header = rows[0]
    loc = {}
    for row in rows[1:]:
        if row is None or row[0] is None:
            continue
        d = dict(zip(header, row))
        key = _norm(d.get("Key"))
        if key:
            loc[key] = {"ko": _norm(d.get("ko")), "en": _norm(d.get("en"))}
    return loc


def localize(loc_map, key, context=""):
    key = _norm(key)
    entry = loc_map.get(key)
    if entry is None:
        print(f"경고: Localization 시트에 없는 키 '{key}'{context} (ko/en 자리에 키 값이 그대로 들어감)")
        return {"ko": key, "en": key}
    return entry


def load_box_sheet(wb, loc_map):
    if "Box" not in wb.sheetnames:
        raise SystemExit("오류: 'Box' 시트가 없습니다.")
    boxes = {}
    for row in sheet_rows(wb["Box"]):
        box_id = _norm(row.get("BoxId"))
        if not box_id:
            continue
        boxes[box_id] = {
            "displayName": localize(loc_map, row.get("NameKey"), f" (Box: {box_id})"),
            "description": localize(loc_map, row.get("DescKey"), f" (Box: {box_id})"),
        }
    return boxes


def load_rarity_pool(wb, valid_rarity_ids):
    if "RarityPool" not in wb.sheetnames:
        raise SystemExit("오류: 'RarityPool' 시트가 없습니다.")
    result = {}
    for row in sheet_rows(wb["RarityPool"]):
        box_id = _norm(row.get("BoxId"))
        rarity_id = _norm(row.get("RarityId"))
        weight = _norm(row.get("Weight"))
        if not box_id or not rarity_id:
            continue
        if rarity_id not in valid_rarity_ids:
            raise SystemExit(
                f"오류: RarityPool의 RarityId '{rarity_id}' (Box: {box_id})가 "
                f"items.json에 존재하는 등급 목록 {sorted(valid_rarity_ids)}에 없습니다."
            )
        result.setdefault(box_id, {})[rarity_id] = float(weight or 0)
    return result


def load_item_pool(ws, box_id, items_by_index):
    pool_by_rarity = {}
    preview_rows = []
    for row in sheet_rows(ws):
        item_index = _norm(row.get("ItemIndex"))
        if not item_index:
            continue
        is_active = _is_true(row.get("IsActive"))
        weight = float(_norm(row.get("Weight")) or 0)

        if item_index not in items_by_index:
            raise SystemExit(
                f"오류: '{box_id}' 시트의 ItemIndex '{item_index}' 가 items.json에 존재하지 않습니다."
            )
        if not is_active:
            continue

        info = items_by_index[item_index]
        rarity_id = info["rarityId"]
        category = info["category"]

        pool_by_rarity.setdefault(rarity_id, []).append({
            "category": category,
            "index": item_index,
            "weight": weight,
        })
        preview_rows.append({
            "box": box_id,
            "index": item_index,
            "rarity": rarity_id,
            "category": category,
            "weight": weight,
        })
    return pool_by_rarity, preview_rows


def compute_final_probabilities(box_id, rarity_weights, pool_by_rarity, preview_rows):
    active_rarities = [r for r in pool_by_rarity if pool_by_rarity[r]]
    total_rarity_weight = sum(rarity_weights.get(r, 0.0) for r in active_rarities)

    for r in active_rarities:
        if r not in rarity_weights:
            print(f"경고: '{box_id}'의 아이템 중 등급 '{r}'이 있는데 RarityPool에 이 등급 가중치가 없습니다 (절대 안 뽑힘).")

    weight_sum_by_rarity = {
        r: sum(it["weight"] for it in items) for r, items in pool_by_rarity.items()
    }

    for row in preview_rows:
        r = row["rarity"]
        rarity_share = (rarity_weights.get(r, 0.0) / total_rarity_weight) if total_rarity_weight > 0 else 0.0
        item_share = (row["weight"] / weight_sum_by_rarity[r]) if weight_sum_by_rarity[r] > 0 else 0.0
        row["finalPercent"] = rarity_share * item_share * 100.0


def main():
    if len(sys.argv) < 3:
        print("사용법: python RandomBoxConverter.py <RandomBox.xlsx> <출력폴더>")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    items_by_index = load_items_json(out_dir)
    valid_rarity_ids = {info["rarityId"] for info in items_by_index.values() if info["rarityId"]}

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    loc_map = load_localization(wb)
    boxes = load_box_sheet(wb, loc_map)
    rarity_pools = load_rarity_pool(wb, valid_rarity_ids)

    result_boxes = {}
    all_preview_rows = []

    for box_id, box_meta in boxes.items():
        if box_id not in wb.sheetnames:
            raise SystemExit(f"오류: Box 시트에 '{box_id}'가 있는데, 그 이름의 ItemPool 시트가 없습니다.")

        rarity_weights = rarity_pools.get(box_id, {})
        pool_by_rarity, preview_rows = load_item_pool(wb[box_id], box_id, items_by_index)

        compute_final_probabilities(box_id, rarity_weights, pool_by_rarity, preview_rows)
        all_preview_rows.extend(preview_rows)

        result_boxes[box_id] = {
            "displayName": box_meta["displayName"],
            "description": box_meta["description"],
            "rarityWeights": rarity_weights,
            "pools": {
                rarity_id: {"items": items}
                for rarity_id, items in pool_by_rarity.items()
            },
        }

    result = {"boxes": result_boxes}

    out_json_path = out_dir / "gacha_tables.json"
    with open(out_json_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=1)
    print(f"{len(result_boxes)}개 박스 저장 완료 -> {out_json_path}")

    print()
    print("[최종 확률 미리보기]")
    for box_id in result_boxes:
        print(f"-- {box_id} --")
        print(f"{'ItemIndex':<12}{'Rarity':<10}{'Category':<10}{'Weight':>8}{'최종확률':>10}")
        for row in sorted(
            (r for r in all_preview_rows if r["box"] == box_id),
            key=lambda r: -r["finalPercent"],
        ):
            print(
                f"{row['index']:<12}{row['rarity']:<10}{row['category']:<10}"
                f"{row['weight']:>8.1f}{row['finalPercent']:>9.2f}%"
            )

    preview_csv_path = out_dir / "gacha_probability_preview.csv"
    with open(preview_csv_path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(["BoxId", "ItemIndex", "RarityId", "Category", "Weight", "FinalProbability(%)"])
        for row in all_preview_rows:
            writer.writerow([
                row["box"], row["index"], row["rarity"], row["category"],
                row["weight"], f"{row['finalPercent']:.4f}",
            ])
    print()
    print(f"확률 미리보기 CSV 저장 완료 -> {preview_csv_path}")


if __name__ == "__main__":
    main()
