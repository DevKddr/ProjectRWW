"""
items.xlsx -> items.json 컨버터.

WeaponDataConverter.py와 같은 패턴을 쓴다:
  - 카테고리별 시트("Weapon", "Item" 등 - "Localization"/"Legend"를 뺀 나머지 전부):
    1행=헤더, 2행=타입, 3행부터 데이터. 시트를 몇 개를 두든, 이름을 뭘로 하든 상관없다 -
    "Localization"/"Legend"가 아니면 전부 아이템 카테고리 시트로 취급해서 읽는다.
    (엑셀에서는 카테고리별로 나뉘어 있지만, Category 컬럼 값 자체는 각 행에 그대로
    남아있다 - 시트 이름과 항상 일치해야 한다.)
  - "Localization" 시트: 1행=헤더(Key/ko), 2행부터 데이터 (타입행 없음).
    언리얼 Localization Dashboard 방식으로 바뀌어서, 이 시트의 문장은 더 이상 items.json에
    들어가지 않는다. items.json에는 키(nameKey/descriptionKey)만 들어가고, 실제 문장은
    아래 StringTable CSV로 나간다.
    - 소스 언어(원문)는 ko. 영어 등 번역문은 이 시트에 두지 않고 Localization Dashboard에서만
      관리한다(엑셀에는 Key와 ko만 있다).

출력 (output_dir 안):
    items.json        : JSON 배열. name/description 대신 nameKey/descriptionKey(키 문자열).
                        FJsonObjectConverter::JsonArrayStringToUStruct로 파싱 가능한 형태.
    ST_Items.csv      : StringTable 임포트용(Key,SourceString). 소스 언어(ko) 문장.
                        언리얼에서 /Game/Core/Text/ST_Items 에 임포트한다.

사용법:
    python ItemDataConverter.py items.xlsx ./output
"""
import sys
import json
from pathlib import Path
import openpyxl

# 아이템 카테고리 시트가 아닌, 고정된 용도의 시트 이름들.
NON_CATEGORY_SHEETS = {"Localization", "Legend"}


def _normalize_value(v):
    """엑셀 드롭다운 검증에 쓴 'TRUE'/'FALSE' 문자열을 실제 JSON bool로 변환."""
    if isinstance(v, str) and v in ("TRUE", "FALSE"):
        return v == "TRUE"
    return v


def sheet_rows(ws, has_type_row=True):
    """1행=헤더, (있다면) 2행=타입, 그 다음부터 데이터. 헤더:값 dict 리스트로 반환."""
    rows = list(ws.iter_rows(values_only=True))
    min_rows = 3 if has_type_row else 2
    if len(rows) < min_rows:
        return []
    headers = rows[0]
    data_start = 2 if has_type_row else 1
    out = []
    for row in rows[data_start:]:
        if row[0] is None:
            continue
        out.append({h: _normalize_value(v) for h, v in zip(headers, row)})
    return out


def load_items(xlsx_path):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    localization = {}
    for row in sheet_rows(wb["Localization"], has_type_row=False):
        localization[row["Key"]] = row["ko"]

    used_keys = []

    def use_key(key):
        """키를 그대로 JSON에 넣고, 사용된 키를 기록한다. 시트에 없는 키는 경고만 한다."""
        if key not in localization:
            print(f"[WARN] Localization 시트에 없는 키: {key} (StringTable에도 들어가지 않는다)")
        elif key not in used_keys:
            used_keys.append(key)
        return key

    items = []
    for sheet_name in wb.sheetnames:
        if sheet_name in NON_CATEGORY_SHEETS:
            continue
        for row in sheet_rows(wb[sheet_name]):
            if row["Category"] != sheet_name:
                print(f"[WARN] [{sheet_name}] {row['Index']}: Category 값('{row['Category']}')이 "
                      f"시트 이름('{sheet_name}')과 다릅니다. 잘못된 시트에 들어있는 건 아닌지 확인하세요.")
            icon_path = row.get("IconPath") or ""
            hud_icon_path = row.get("HudIconPath") or ""
            actor_class_path = row.get("ActorClassPath") or ""
            items.append({
                "index": row["Index"],
                "category": row["Category"],
                "nameKey": use_key(row["NameKey"]),
                "descriptionKey": use_key(row["DescKey"]),
                "rarityId": row["RarityId"],
                "iconPath": icon_path,
                "hudIconPath": hud_icon_path,
                "actorClassPath": actor_class_path,
                "equipTime": float(row["EquipTime"]),
                "primarySkillCooldown": float(row["PrimarySkillCooldown"]),
                "primarySkillManaCost": float(row["PrimarySkillManaCost"]),
                "primarySkillCastTime": float(row["PrimarySkillCastTime"]),
                "secondarySkillCooldown": float(row["SecondarySkillCooldown"]),
                "secondarySkillManaCost": float(row["SecondarySkillManaCost"]),
                "secondarySkillCastTime": float(row["SecondarySkillCastTime"]),
            })

    # 사용된 키만, 시트 순서대로 StringTable용 문장 목록을 만든다.
    strings = [(key, localization[key]) for key in used_keys]
    return items, strings


def write_string_table_csv(path, rows):
    """언리얼 StringTable CSV 형식(Key,SourceString). UTF-8 BOM은 붙이지 않는다."""
    import csv
    with open(path, "w", encoding="utf-8", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(["Key", "SourceString"])
        for key, text in rows:
            writer.writerow([key, text])


def main():
    if len(sys.argv) != 3:
        print("usage: python ItemDataConverter.py <items.xlsx> <output_dir>")
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    out_dir = Path(sys.argv[2])
    out_dir.mkdir(parents=True, exist_ok=True)

    items, strings = load_items(xlsx_path)

    out_path = out_dir / "items.json"
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)

    ko_path = out_dir / "ST_Items.csv"
    write_string_table_csv(ko_path, strings)

    print(f"[OK] {out_path}  (아이템 {len(items)}개)")
    print(f"[OK] {ko_path}  (소스 문장 {len(strings)}개)")


if __name__ == "__main__":
    main()
